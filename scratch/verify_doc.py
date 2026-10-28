import docx

doc = docx.Document('LogShield_Final_Report.docx')
text = ' '.join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            text += ' ' + cell.text

em_dash_count = text.count('—')
en_dash_count = text.count('–')
print(f'Em dashes (—): {em_dash_count}')
print(f'En dashes (–): {en_dash_count}')

slop_words = ['delve', 'tapestry', 'testament', 'beacon', 'crucial role', 'furthermore, it is important to note', 'in conclusion, it is evident']
for w in slop_words:
    c = text.lower().count(w)
    print(f'Slop count "{w}": {c}')

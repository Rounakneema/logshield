import docx

d = docx.Document('LogShield_Final_Report_backup.docx')

for i, p in enumerate(d.paragraphs):
    if 'drawing' in p._p.xml or 'graphic' in p._p.xml:
        txt = p.text.strip()
        prev = d.paragraphs[i-1].text.strip() if i > 0 else ''
        nxt = d.paragraphs[i+1].text.strip() if i+1 < len(d.paragraphs) else ''
        print(f'P{i:3d}: txt="{txt}" | prev="{prev[:60]}" | nxt="{nxt[:60]}"')

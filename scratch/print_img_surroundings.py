import docx

d = docx.Document('LogShield_Final_Report_backup.docx')

img_idx = 1
for i, p in enumerate(d.paragraphs):
    if 'drawing' in p._p.xml or 'graphic' in p._p.xml:
        prev = ''
        for k in range(i-1, -1, -1):
            if d.paragraphs[k].text.strip():
                prev = d.paragraphs[k].text.strip()
                break
        nxt = ''
        for k in range(i+1, len(d.paragraphs)):
            if d.paragraphs[k].text.strip():
                nxt = d.paragraphs[k].text.strip()
                break
        print(f'Img {img_idx:2d} (P{i:3d}):')
        print(f'   Prev: "{prev[:70]}"')
        print(f'   Nxt:  "{nxt[:70]}"')
        img_idx += 1

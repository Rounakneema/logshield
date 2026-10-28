import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

doc = docx.Document()
sec1 = doc.sections[0]
sec1.top_margin = Inches(0.75)
sec1.bottom_margin = Inches(1.0)
sec1.left_margin = Inches(0.625)
sec1.right_margin = Inches(0.625)

p = doc.add_paragraph("Title in 1 column")

# 2-column section
sec2 = doc.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
sec2.top_margin = Inches(0.75); sec2.bottom_margin = Inches(1.0)
sec2.left_margin = Inches(0.625); sec2.right_margin = Inches(0.625)
sec2._sectPr.append(parse_xml(r'<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>'))

p2 = doc.add_paragraph("This is column 1 text in section 2.")

# Full-width table section
sec3 = doc.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
sec3.top_margin = Inches(0.75); sec3.bottom_margin = Inches(1.0)
sec3.left_margin = Inches(0.625); sec3.right_margin = Inches(0.625)
sec3._sectPr.append(parse_xml(r'<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="1"/>'))

p_tbl_title = doc.add_paragraph("TABLE IV: Full Width Table")
tbl = doc.add_table(rows=3, cols=8)
tbl.autofit = True

# Resume 2 columns
sec4 = doc.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
sec4.top_margin = Inches(0.75); sec4.bottom_margin = Inches(1.0)
sec4.left_margin = Inches(0.625); sec4.right_margin = Inches(0.625)
sec4._sectPr.append(parse_xml(r'<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>'))

p3 = doc.add_paragraph("This is column text resuming in 2 columns after the table.")

doc.save("scratch/test_section_breaks.docx")
print("Saved test_section_breaks.docx")

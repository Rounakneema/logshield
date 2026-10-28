import os, sys, re, zipfile, io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sort_references import old_to_new, sorted_refs
from remap_citations import remap_text_citations

def create_element(name):
    return OxmlElement(name)

def set_cell_border(cell, **kwargs):
    """
    Set cell borders
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='000000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))

def set_ieee_table_style(table):
    """Applies official IEEE 3-line booktabs table style."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Top rule 1 pt (sz=8 is 1 pt in eighths of pt? In Word sz=4 is 0.5pt, sz=8 is 1pt)
    top_border = dict(sz=8, val='single', color='000000')
    header_bot = dict(sz=4, val='single', color='000000')
    bot_border = dict(sz=8, val='single', color='000000')
    none_border = dict(val='none')
    
    num_rows = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            # Clear existing vertical borders
            borders = {
                'left': none_border,
                'right': none_border,
                'insideV': none_border,
            }
            if r_idx == 0:
                borders['top'] = top_border
                borders['bottom'] = header_bot
            elif r_idx == num_rows - 1:
                borders['top'] = none_border
                borders['bottom'] = bot_border
            else:
                borders['top'] = none_border
                borders['bottom'] = none_border
            set_cell_border(cell, **borders)

def build_official_ieee_docx():
    src_docx = r'd:\LOGSHEILD\LogShield_IEEE_Paper_Updated.docx'
    out_docx = r'd:\LOGSHEILD\LogShield_IEEE_Paper_Official.docx'
    
    doc_src = docx.Document(src_docx)
    doc_out = docx.Document()
    
    # ── Section 1: Title & Authors (1 Column) ──────────────────────────────────
    sec1 = doc_out.sections[0]
    sec1.page_width = Inches(8.5)
    sec1.page_height = Inches(11.0)
    sec1.top_margin = Inches(0.75)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(0.625)
    sec1.right_margin = Inches(0.625)
    
    # 1. Paper Title
    p_title = doc_out.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.05
    run_title = p_title.add_run("LogShield: In-Pod Secret Sanitization and Autonomous Closed-Loop Containment for Kubernetes Log Streams")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    
    # 2. Authors Table (2 Columns across top)
    t_auth = doc_out.add_table(rows=1, cols=2)
    t_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_auth.autofit = False
    
    c1, c2 = t_auth.rows[0].cells
    c1.width = Inches(3.5)
    c2.width = Inches(3.5)
    
    p_a1 = c1.paragraphs[0]
    p_a1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a1.paragraph_format.line_spacing = 1.05
    p_a1.paragraph_format.space_after = Pt(2)
    r = p_a1.add_run("Rounak Neema\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(11); r.font.bold = True
    r = p_a1.add_run("Department of Computer Engineering\nMukesh Patel School of Technology\nManagement & Engineering\nNMIMS University\nMumbai, India\nrounak.neema@nmims.edu")
    r.font.name = "Times New Roman"; r.font.size = Pt(9.5)
    
    p_a2 = c2.paragraphs[0]
    p_a2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a2.paragraph_format.line_spacing = 1.05
    p_a2.paragraph_format.space_after = Pt(2)
    r = p_a2.add_run("Capstone Research Group\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(11); r.font.bold = True
    r = p_a2.add_run("Cloud-Native Security Laboratory Project\nLogShield Initiative\nMumbai, India\nresearch@logshield.internal\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(9.5)
    
    # Remove borders from author table
    for cell in [c1, c2]:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                              r'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
                              r'</w:tcBorders>')
        tcPr.append(tcBorders)
    
    # ── Section 2: Body (2 Columns) ──────────────────────────────────────────
    sec2 = doc_out.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
    sec2.top_margin = Inches(0.75)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(0.625)
    sec2.right_margin = Inches(0.625)
    
    # Enable 2-column layout in Section 2
    sectPr = sec2._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), '360')  # 0.25 inch = 18 pt = 360 twips
    else:
        cols_xml = parse_xml(r'<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>')
        sectPr.append(cols_xml)
        
    print("Set up Section 1 (1 column) and Section 2 (2 columns)")
    
    # Map of images
    image_paths = {
        "fig1": r"d:\LOGSHEILD\paper\figures\fig1_architecture.png",
        "fig2": r"d:\LOGSHEILD\paper\figures\fig2_pipeline.png",
        "fig3": r"d:\LOGSHEILD\paper\figures\fig6_lineage.png",
        "fig4": r"d:\LOGSHEILD\paper\figures\fig4_ablation.png",
        "fig5": r"d:\LOGSHEILD\paper\figures\fig7_evasion_benchmark.png",
        "fig6": r"d:\LOGSHEILD\paper\figures\fig8_containment_lifecycle.png",
        "fig7": r"d:\LOGSHEILD\paper\figures\fig5_benchmark.png",
    }
    
    # Tables from doc_src:
    # Table 0: Author table in src (skip)
    # Table 1: Code snippet in Section I
    # Table 2: TABLE I (Mathematical Notations)
    # Table 3: Code snippet in Section IV-A
    # Table 4: TABLE II (Feature Space and Parameter Weights)
    # Table 5: Algorithm 1
    # Table 6: Algorithm 2
    # Table 7: TABLE III (Authoritative Inventory)
    # Table 8: TABLE IV (Ablation Study)
    # Table 9: TABLE V (Adversarial Evasion)
    # Table 10: TABLE VI (Multi-Baseline Comparison)
    # Table 11: TABLE VII (Latency and Throughput)
    # Table 12: TABLE VIII (Cryptographic Profile)

    # Process all paragraphs from doc_src (starting from Abstract at index 3)
    skip_indices = {0, 1, 2} # Skip title and authors from src
    
    p_idx = 3
    total_p = len(doc_src.paragraphs)
    table_cursor = 1 # Start from table 1
    
    while p_idx < total_p:
        # Check if we hit References section
        p_raw = doc_src.paragraphs[p_idx]
        text_raw = p_raw.text.strip()
        
        if text_raw == "References" or text_raw.startswith("References"):
            break # Stop at References to handle separately
            
        # Check if paragraph is Abstract
        if text_raw.startswith("Abstract"):
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_out.paragraph_format.line_spacing = 1.0
            p_out.paragraph_format.space_after = Pt(4)
            p_out.paragraph_format.first_line_indent = Pt(0)
            
            # Extract content after Abstract—
            m = re.match(r'Abstract[—\-\s]+(.*)', text_raw, re.DOTALL)
            body = m.group(1) if m else text_raw
            body = remap_text_citations(body)
            
            r_head = p_out.add_run("Abstract—")
            r_head.font.name = "Times New Roman"
            r_head.font.size = Pt(9.0)
            r_head.font.bold = True
            r_head.font.italic = True
            
            r_body = p_out.add_run(body)
            r_body.font.name = "Times New Roman"
            r_body.font.size = Pt(9.0)
            r_body.font.bold = True
            p_idx += 1
            continue

        # Check if paragraph is Index Terms
        if text_raw.startswith("Index Terms"):
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_out.paragraph_format.line_spacing = 1.0
            p_out.paragraph_format.space_after = Pt(6)
            p_out.paragraph_format.first_line_indent = Pt(0)
            
            m = re.match(r'Index Terms[—\-\s]+(.*)', text_raw, re.DOTALL)
            body = m.group(1) if m else text_raw
            
            r_head = p_out.add_run("Index Terms—")
            r_head.font.name = "Times New Roman"
            r_head.font.size = Pt(9.0)
            r_head.font.bold = True
            r_head.font.italic = True
            
            r_body = p_out.add_run(body)
            r_body.font.name = "Times New Roman"
            r_body.font.size = Pt(9.0)
            r_body.font.bold = True
            p_idx += 1
            continue
            
        if text_raw.startswith("This study was conducted as part of the Capstone"):
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_out.paragraph_format.line_spacing = 1.0
            p_out.paragraph_format.space_after = Pt(8)
            p_out.paragraph_format.first_line_indent = Pt(0)
            r = p_out.add_run(text_raw)
            r.font.name = "Times New Roman"; r.font.size = Pt(8.0); r.font.italic = True
            p_idx += 1
            continue

        # Check for Major Heading Level 1 (e.g. "I. Introduction", "II. Threat Model...")
        m_h1 = re.match(r'^([I|V|X]+\.\s+.*)', text_raw)
        if m_h1:
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_out.paragraph_format.space_before = Pt(10)
            p_out.paragraph_format.space_after = Pt(4)
            p_out.paragraph_format.keep_with_next = True
            p_out.paragraph_format.first_line_indent = Pt(0)
            r = p_out.add_run(text_raw.upper())
            r.font.name = "Times New Roman"; r.font.size = Pt(10.0); r.font.bold = True
            p_idx += 1
            continue

        # Check for Secondary Heading Level 2 (e.g. "A. System Model...", "B. Contributions")
        m_h2 = re.match(r'^([A-Z]\.\s+.*)', text_raw)
        if m_h2 and len(text_raw) < 70 and not text_raw.startswith("TABLE"):
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_out.paragraph_format.space_before = Pt(8)
            p_out.paragraph_format.space_after = Pt(2)
            p_out.paragraph_format.keep_with_next = True
            p_out.paragraph_format.first_line_indent = Pt(0)
            r = p_out.add_run(text_raw)
            r.font.name = "Times New Roman"; r.font.size = Pt(10.0); r.font.italic = True
            p_idx += 1
            continue

        # Check for Acknowledgment
        if text_raw.lower() == "acknowledgment":
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_out.paragraph_format.space_before = Pt(10)
            p_out.paragraph_format.space_after = Pt(4)
            p_out.paragraph_format.keep_with_next = True
            p_out.paragraph_format.first_line_indent = Pt(0)
            r = p_out.add_run("ACKNOWLEDGMENT")
            r.font.name = "Times New Roman"; r.font.size = Pt(10.0); r.font.bold = True
            p_idx += 1
            continue

        # Check for Figure insertion and caption
        if text_raw.startswith("Fig."):
            fig_match = re.match(r'Fig\.\s*(\d+)\.', text_raw)
            fig_num = fig_match.group(1) if fig_match else "1"
            img_key = f"fig{fig_num}"
            
            # Insert image paragraph
            p_img = doc_out.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.first_line_indent = Pt(0)
            
            if img_key in image_paths and os.path.exists(image_paths[img_key]):
                run_img = p_img.add_run()
                # 3.4 inches fits cleanly in a 3.5 inch column!
                run_img.add_picture(image_paths[img_key], width=Inches(3.40))
            
            # Caption paragraph
            p_cap = doc_out.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.paragraph_format.first_line_indent = Pt(0)
            r_cap = p_cap.add_run(text_raw)
            r_cap.font.name = "Times New Roman"; r_cap.font.size = Pt(8.0)
            
            p_idx += 1
            continue

        # Check for Table titles (e.g. TABLE I, TABLE IV...)
        if text_raw.startswith("TABLE "):
            p_tnum = doc_out.add_paragraph()
            p_tnum.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_tnum.paragraph_format.space_before = Pt(8)
            p_tnum.paragraph_format.space_after = Pt(1)
            p_tnum.paragraph_format.keep_with_next = True
            p_tnum.paragraph_format.first_line_indent = Pt(0)
            r = p_tnum.add_run(text_raw)
            r.font.name = "Times New Roman"; r.font.size = Pt(8.5); r.font.bold = True
            
            # Next paragraph is typically table title
            p_idx += 1
            if p_idx < total_p and not doc_src.paragraphs[p_idx].text.strip().startswith(('I.', 'II.', 'A.', 'B.', 'TABLE', 'Fig.')):
                p_ttitle = doc_out.add_paragraph()
                p_ttitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ttitle.paragraph_format.space_before = Pt(1)
                p_ttitle.paragraph_format.space_after = Pt(3)
                p_ttitle.paragraph_format.keep_with_next = True
                p_ttitle.paragraph_format.first_line_indent = Pt(0)
                r_title = p_ttitle.add_run(doc_src.paragraphs[p_idx].text.strip())
                r_title.font.name = "Times New Roman"; r_title.font.size = Pt(8.0); r_title.font.italic = True
                p_idx += 1
            
            # Insert the corresponding table from doc_src
            if table_cursor < len(doc_src.tables):
                src_tbl = doc_src.tables[table_cursor]
                table_cursor += 1
                
                # Clone table into doc_out
                rows_cnt = len(src_tbl.rows)
                cols_cnt = len(src_tbl.columns)
                new_tbl = doc_out.add_table(rows=rows_cnt, cols=cols_cnt)
                new_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                new_tbl.autofit = True
                
                for ri in range(rows_cnt):
                    for ci in range(cols_cnt):
                        txt = src_tbl.rows[ri].cells[ci].text.strip()
                        c_out = new_tbl.rows[ri].cells[ci]
                        c_out.text = txt
                        p_c = c_out.paragraphs[0]
                        p_c.paragraph_format.space_before = Pt(1.5)
                        p_c.paragraph_format.space_after = Pt(1.5)
                        p_c.paragraph_format.line_spacing = 1.0
                        if len(p_c.runs) > 0:
                            p_c.runs[0].font.name = "Times New Roman"
                            p_c.runs[0].font.size = Pt(7.5)
                            if ri == 0 or 'Part' in txt:
                                p_c.runs[0].font.bold = True
                
                set_ieee_table_style(new_tbl)
                
                # Add spacing after table
                p_post = doc_out.add_paragraph()
                p_post.paragraph_format.space_before = Pt(0)
                p_post.paragraph_format.space_after = Pt(6)
            continue

        # Standard Body Paragraph
        if text_raw:
            remapped_text = remap_text_citations(text_raw)
            p_out = doc_out.add_paragraph()
            p_out.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_out.paragraph_format.line_spacing = 1.0
            p_out.paragraph_format.space_before = Pt(0)
            p_out.paragraph_format.space_after = Pt(0)
            p_out.paragraph_format.first_line_indent = Pt(10.8) # 0.15 inch IEEE indent
            
            # Handle run-in italic titles (e.g. "Static Code Scanners:", "Invariant 1:", "Theorem 1:")
            m_runin = re.match(r'^((?:[A-Za-z0-9\-\s\(\)]+[:—])|(?:Invariant \d+[^:]*:)|(?:Theorem \d+[^:]*:))(.*)', remapped_text, re.DOTALL)
            if m_runin and len(m_runin.group(1)) < 60:
                head_str = m_runin.group(1)
                tail_str = m_runin.group(2)
                r_h = p_out.add_run(head_str)
                r_h.font.name = "Times New Roman"; r_h.font.size = Pt(10.0); r_h.font.italic = True
                r_t = p_out.add_run(tail_str)
                r_t.font.name = "Times New Roman"; r_t.font.size = Pt(10.0)
            else:
                r_t = p_out.add_run(remapped_text)
                r_t.font.name = "Times New Roman"; r_t.font.size = Pt(10.0)
                
        p_idx += 1

    # ── References Section ───────────────────────────────────────────────────
    p_ref_head = doc_out.add_paragraph()
    p_ref_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ref_head.paragraph_format.space_before = Pt(12)
    p_ref_head.paragraph_format.space_after = Pt(6)
    p_ref_head.paragraph_format.keep_with_next = True
    p_ref_head.paragraph_format.first_line_indent = Pt(0)
    r_rf = p_ref_head.add_run("REFERENCES")
    r_rf.font.name = "Times New Roman"; r_rf.font.size = Pt(10.0); r_rf.font.bold = True
    
    # Add each sorted reference (Recent UP, Older DOWN)
    for idx, ref in enumerate(sorted_refs, 1):
        p_ref = doc_out.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.line_spacing = 1.0
        p_ref.paragraph_format.space_before = Pt(1.5)
        p_ref.paragraph_format.space_after = Pt(1.5)
        p_ref.paragraph_format.left_indent = Inches(0.25)
        p_ref.paragraph_format.first_line_indent = Inches(-0.25) # Hanging indent
        
        # Add [N]
        r_num = p_ref.add_run(f"[{idx}]\t")
        r_num.font.name = "Times New Roman"; r_num.font.size = Pt(8.0)
        
        # Add citation text
        r_txt = p_ref.add_run(ref["citation"])
        r_txt.font.name = "Times New Roman"; r_txt.font.size = Pt(8.0)

    doc_out.save(out_docx)
    print(f"Successfully generated: {out_docx} ({os.path.getsize(out_docx)} bytes)")

if __name__ == "__main__":
    build_official_ieee_docx()

import docx

doc = docx.Document('LogShield_IEEE_Paper_Official.docx')
print('Sections:', len(doc.sections))
print('Paragraphs:', len(doc.paragraphs))
print('Tables:', len(doc.tables))
for i, s in enumerate(doc.sections):
    cols = s._sectPr.xpath('./w:cols')
    num = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num') if cols else '1'
    print(f'Sec {i+1}: num_cols={num}')

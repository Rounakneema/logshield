import docx
import zipfile
import xml.etree.ElementTree as ET

doc = docx.Document('LogShield_Final_Report_v7_Complete.docx')

print('=== 1. PARAGRAPH COUNT & SECTIONS ===')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')

print('\n=== 2. HEADINGS CHECK ===')
h1_list = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1' or p.text.startswith('CHAPTER') or p.text.startswith('APPENDIX')]
for h in h1_list:
    print(' ', h)

print('\n=== 3. FIGURE & IMAGE MAPPING VERIFICATION ===')
with zipfile.ZipFile('LogShield_Final_Report_v7_Complete.docx') as z:
    rels_xml = z.read('word/_rels/document.xml.rels')
    rels = ET.fromstring(rels_xml)
    rId_to_target = {rel.attrib['Id']: rel.attrib['Target'] for rel in rels}

    doc_xml = z.read('word/document.xml')
    tree = ET.fromstring(doc_xml)
    
    current_img = None
    fig_to_img = {}
    for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
        text = ''.join(p.itertext()).strip()
        blips = [el.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') 
                 for el in p.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip')]
        blips = [b for b in blips if b]
        if blips:
            current_img = [rId_to_target.get(b) for b in blips]
        if text.startswith('Figure ') and ':' in text and not text.endswith('\t'):
            # Body figure caption
            fig_num = text.split(':')[0].replace('Figure ', '').strip()
            if fig_num not in fig_to_img and current_img:
                fig_to_img[fig_num] = (current_img[0], text[:65])

print(f'Unique Body Figures Detected: {len(fig_to_img)}')
img_counts = {}
for num, (target, cap) in sorted(fig_to_img.items()):
    img_counts[target] = img_counts.get(target, 0) + 1
    print(f'Fig {num:4s} -> Image: {target:35s} | {cap}')

duplicates = {target: count for target, count in img_counts.items() if count > 1}
if duplicates:
    print('\n[WARNING] Duplicate Images Found across figures:', duplicates)
else:
    print('\n[SUCCESS] ZERO DUPLICATE IMAGES! Every figure has a unique distinct image!')

print('\n=== 4. APPENDIX A CODE BLOCKS VERIFICATION ===')
found_listings = []
for p in doc.paragraphs:
    if 'Listing ' in p.text:
        found_listings.append(p.text)

print(f'Listings found: {len(found_listings)}')
for l in found_listings:
    print(' ', l)

print('\n=== 5. CHECK FOR EM DASHES & EN DASHES ===')
dash_count = 0
for p in doc.paragraphs:
    if '—' in p.text or '–' in p.text:
        dash_count += 1
print(f'Paragraphs with em/en dashes: {dash_count} (Must be 0)')

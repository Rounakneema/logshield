import openpyxl
import json
from pathlib import Path
import os

def import_secretbench(xlsx_path: str, output_dir: str):
    print(f"Loading {xlsx_path}...")
    wb = openpyxl.load_workbook(xlsx_path)
    
    # --- Sheet 1: Full catalog ---
    ws1 = wb['Sheet1']
    patterns = []
    for row in ws1.iter_rows(min_row=2, values_only=True):
        if row[1] and row[2]:
            pattern_id = row[0]
            # Handle '=ROW(A1)' format
            if isinstance(pattern_id, str) and pattern_id.startswith('=ROW(A'):
                try:
                    pattern_id = int(pattern_id.replace('=ROW(A', '').replace(')', ''))
                except ValueError:
                    pass
            patterns.append({
                "id": pattern_id,
                "name": row[1],
                "regex": row[2],
                "source": row[3],
                "fp_risk": "LOW"   # default
            })
    print(f"Loaded {len(patterns)} patterns from Sheet1")

    # --- Sheet 2: High-FP denylist ---
    ws2 = wb['Sheet2']
    high_fp_ids = set()
    high_fp_count = 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row[0]:
            try:
                row_id = int(float(row[0]))
                high_fp_ids.add(row_id)
                # Override fp_risk
                for p in patterns:
                    if p['id'] == row_id:
                        p['fp_risk'] = 'HIGH'
                        high_fp_count += 1
            except ValueError:
                pass
    
    print(f"Loaded {len(high_fp_ids)} high-FP IDs from Sheet2")
    print(f"Marked {high_fp_count} patterns as HIGH risk")

    result = {"patterns": patterns, "high_fp_ids": list(high_fp_ids)}

    os.makedirs(output_dir, exist_ok=True)
    output_path = Path(output_dir) / "secretbench_761.json"
    
    with open(output_path, 'w') as f:
        json.dump(result, f, indent=2)
    print(f"Exported {len(result['patterns'])} patterns to {output_path}")

if __name__ == "__main__":
    import_secretbench(
        r'd:\LOGSHEILD\Secret Regular Expression (1).xlsx',
        r'd:\LOGSHEILD\sidecar\logmask\catalog'
    )
